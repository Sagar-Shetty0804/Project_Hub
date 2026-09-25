from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.text import slugify
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.permissions import role_required
from notifications.utils import notify
from projects.models import Project

from . import services
from .forms import ReviewRoundForm, RubricForm
from .models import Criterion, Remark, ReviewRound, Rubric, Score


@login_required
def index(request):
    user = request.user
    if user.is_student:
        return render(request, "evaluation/my_marks.html", {"results": services.student_results(user)})
    if not user.is_faculty:
        raise PermissionDenied
    rounds = list(ReviewRound.objects.select_related("rubric"))
    for r in rounds:
        r.done, r.total = services.progress(r)
        r.pct = round(r.done * 100 / r.total) if r.total else 0
    return render(request, "evaluation/index.html", {
        "rounds": rounds,
        "rubrics": Rubric.objects.prefetch_related("criteria") if user.is_admin_role else None,
    })


# -------------------------------------------------------------- rubrics ----

@role_required()
def rubric_edit(request, pk=None):
    rubric = get_object_or_404(Rubric, pk=pk) if pk else None
    form = RubricForm(request.POST or None, instance=rubric)
    criteria = list(rubric.criteria.all()) if rubric else []

    if request.method == "POST" and form.is_valid():
        names = request.POST.getlist("c_name")
        maxes = request.POST.getlist("c_max")
        descs = request.POST.getlist("c_desc")
        rows = []
        for i, name in enumerate(names):
            name = name.strip()
            if not name:
                continue
            try:
                mx = Decimal(maxes[i])
                assert mx > 0
            except (InvalidOperation, AssertionError, IndexError):
                messages.error(request, f"“{name}” needs a maximum mark above 0.")
                return render(request, "evaluation/rubric_form.html", {"form": form, "rubric": rubric, "criteria": criteria})
            rows.append((name[:60], mx, descs[i][:200] if i < len(descs) else ""))
        if not rows:
            messages.error(request, "Add at least one criterion.")
        elif rubric and rubric.rounds.filter(scores__isnull=False).exists() and len(rows) != len(criteria):
            messages.error(request, "This rubric already has marks entered, so criteria can be renamed but not added or removed.")
        else:
            with transaction.atomic():
                rubric = form.save(commit=False)
                rubric.created_by = rubric.created_by or request.user
                rubric.save()
                existing = list(rubric.criteria.all())
                for i, (name, mx, desc) in enumerate(rows):
                    c = existing[i] if i < len(existing) else Criterion(rubric=rubric)
                    c.name, c.max_marks, c.description, c.order = name, mx, desc, i
                    c.save()
                for extra in existing[len(rows):]:
                    extra.delete()
            messages.success(request, f"Rubric “{rubric.name}” saved ({rubric.total_marks} marks).")
            return redirect("evaluation:index")
    return render(request, "evaluation/rubric_form.html", {"form": form, "rubric": rubric, "criteria": criteria})


# --------------------------------------------------------------- rounds ----

@role_required()
def round_create(request):
    if not Rubric.objects.exists():
        messages.info(request, "Create a rubric first; a review round marks projects against one.")
        return redirect("evaluation:rubric_new")
    form = ReviewRoundForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        rnd = form.save(commit=False)
        rnd.created_by = request.user
        rnd.save()
        messages.success(request, f"“{rnd.name}” is open for marking.")
        return redirect("evaluation:round", pk=rnd.pk)
    return render(request, "evaluation/round_form.html", {"form": form})


@login_required
def round_detail(request, pk):
    user = request.user
    if not user.is_faculty:
        raise PermissionDenied
    rnd = get_object_or_404(ReviewRound.objects.select_related("rubric"), pk=pk)

    if request.method == "POST" and user.is_admin_role:
        action = request.POST.get("action")
        if action == "lock":
            rnd.state = ReviewRound.State.LOCKED if rnd.is_open else ReviewRound.State.OPEN
            rnd.save(update_fields=["state"])
            messages.success(request, "Marks locked." if not rnd.is_open else "Round reopened for marking.")
        elif action == "publish":
            rnd.results_published = not rnd.results_published
            rnd.save(update_fields=["results_published"])
            if rnd.results_published:
                students = {s.student for s in rnd.scores.select_related("student")}
                notify(students, f"Results for “{rnd.name}” are out", reverse("evaluation:index"), actor=user, email=True)
                messages.success(request, "Results published to students.")
            else:
                messages.info(request, "Results hidden from students.")
        elif action == "delete" and not rnd.scores.exists():
            rnd.delete()
            messages.info(request, "Round deleted.")
            return redirect("evaluation:index")
        return redirect("evaluation:round", pk=rnd.pk)

    projects = rnd.projects().select_related("guide").prefetch_related("memberships__student")
    if user.is_guide and not user.is_admin_role:
        projects = projects.filter(guide=user)
    projects = list(projects.order_by("code"))
    criteria, rows = services.round_table(rnd, projects)
    for row in rows:
        row["done"] = sum(s["complete"] for s in row["students"])
        row["can_mark"] = rnd.can_mark(user, row["project"])
    done = sum(r["done"] for r in rows)
    total = sum(len(r["students"]) for r in rows)
    return render(request, "evaluation/round.html", {
        "round": rnd, "criteria": criteria, "rows": rows, "done": done, "total": total,
        "pct": round(done * 100 / total) if total else 0,
    })


@login_required
def mark(request, pk, code):
    rnd = get_object_or_404(ReviewRound.objects.select_related("rubric"), pk=pk)
    project = get_object_or_404(rnd.projects(), code=code)
    user = request.user
    if not (user.is_faculty and project.can_view_private(user)):
        raise PermissionDenied
    editable = rnd.can_mark(user, project)
    criteria, rows = services.round_table(rnd, [project])
    row = rows[0]

    if request.method == "POST":
        if not editable:
            raise PermissionDenied
        errors = []
        with transaction.atomic():
            for s in row["students"]:
                for c in criteria:
                    raw = request.POST.get(f"m_{s['student'].id}_{c.id}", "").strip()
                    if raw == "":
                        Score.objects.filter(round=rnd, student=s["student"], criterion=c).delete()
                        continue
                    try:
                        val = Decimal(raw)
                    except InvalidOperation:
                        errors.append(f"{s['student'].display_name}: “{raw}” isn't a number.")
                        continue
                    if not (0 <= val <= c.max_marks):
                        errors.append(f"{s['student'].display_name}: {c.name} must be between 0 and {c.max_marks}.")
                        continue
                    Score.objects.update_or_create(
                        round=rnd, student=s["student"], criterion=c,
                        defaults={"marks": val, "project": project, "marked_by": user},
                    )
            Remark.objects.update_or_create(
                round=rnd, project=project,
                defaults={"text": request.POST.get("remark", "").strip(), "author": user},
            )
        for e in errors:
            messages.error(request, e)
        if not errors:
            messages.success(request, f"Marks saved for {project.title}.")
            if request.POST.get("next") == "round":
                return redirect("evaluation:round", pk=rnd.pk)
        return redirect("evaluation:mark", pk=rnd.pk, code=project.code)

    # Next project to mark, for a smooth "save & next" flow
    scope = list(rnd.projects().order_by("code").values_list("code", flat=True))
    idx = scope.index(project.code)
    return render(request, "evaluation/mark.html", {
        "round": rnd, "project": project, "criteria": criteria, "row": row, "editable": editable,
        "prev_code": scope[idx - 1] if idx > 0 else None,
        "next_code": scope[idx + 1] if idx + 1 < len(scope) else None,
        "out_of": rnd.rubric.total_marks,
    })


@login_required
def export(request, pk):
    user = request.user
    if not user.is_faculty:
        raise PermissionDenied
    rnd = get_object_or_404(ReviewRound.objects.select_related("rubric"), pk=pk)
    projects = rnd.projects().select_related("guide").prefetch_related("memberships__student")
    if user.is_guide and not user.is_admin_role:
        projects = projects.filter(guide=user)
    criteria, rows = services.round_table(rnd, list(projects.order_by("code")))

    wb = Workbook()
    ws = wb.active
    ws.title = "Marks"
    ws.append([rnd.name])
    ws.append([f"{rnd.scope_label()} · Rubric: {rnd.rubric.name}"])
    ws.append([])
    header = ["Group", "Project", "Guide", "Student", "Class", "Email"]
    header += [f"{c.name} (/{c.max_marks})" for c in criteria] + [f"Total (/{rnd.rubric.total_marks})", "Remarks"]
    ws.append(header)
    head_row = ws.max_row
    for cell in ws[head_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2C4FD6")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws["A1"].font = Font(bold=True, size=14)

    for r in rows:
        p = r["project"]
        for s in r["students"]:
            st = s["student"]
            ws.append(
                [p.code, p.title, p.guide.display_name if p.guide else "", st.display_name, st.class_label, st.email]
                + [float(v) if v is not None else None for v in s["cells"]]
                + [float(s["total"]), r["remark"].text if r["remark"] else ""]
            )
    widths = [16, 34, 20, 22, 16, 28] + [14] * len(criteria) + [12, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=head_row + 1, column=5)

    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{slugify(rnd.name) or "marks"}.xlsx"'
    return resp


@role_required()
@require_POST
def toggle_featured(request, code):
    project = get_object_or_404(Project, code=code)
    project.is_featured = not project.is_featured
    project.save(update_fields=["is_featured"])
    messages.success(request, "Added to featured projects." if project.is_featured else "Removed from featured.")
    return redirect(project)
